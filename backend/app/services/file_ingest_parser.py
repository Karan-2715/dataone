"""
Parse uploaded tabular files into (columns, rows) for Unity Catalog ingestion.

Supported formats: CSV, TSV, JSON (array of objects), JSON Lines, Parquet, Excel (.xlsx).
Every value is returned as ``str`` or ``None`` — the ingest endpoint stores columns as STRING.
"""
import csv
import io
import json
import logging
import os
import re
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".parquet", ".xlsx")

Row = Dict[str, Optional[str]]


class UnsupportedFileError(ValueError):
    """Raised when the file extension or content cannot be parsed."""


def table_name_from_filename(filename: str) -> str:
    """Derive a safe Unity Catalog table name from an uploaded filename."""
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    name = re.sub(r"[^A-Za-z0-9_]+", "_", stem).strip("_").lower()
    if not name:
        name = "uploaded_file"
    if not name[0].isalpha():
        name = "t_" + name
    return name


def _stringify(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str, ensure_ascii=False)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def _unique_columns(raw: List[Any]) -> List[str]:
    """Name blank columns and de-duplicate repeated ones so the CREATE TABLE is valid."""
    seen: Dict[str, int] = {}
    columns: List[str] = []
    for idx, col in enumerate(raw):
        name = (str(col).strip() if col is not None else "") or f"column_{idx + 1}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        columns.append(name if count == 0 else f"{name}_{count + 1}")
    return columns


def _rows_from_dicts(records: List[Dict[str, Any]]) -> Tuple[List[str], List[Row]]:
    keys: List[str] = []
    seen = set()
    for rec in records:
        for key in rec:
            if key not in seen:
                seen.add(key)
                keys.append(key)
    columns = _unique_columns(keys)
    rows = [{columns[i]: _stringify(rec.get(k)) for i, k in enumerate(keys)} for rec in records]
    return columns, rows


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows commonly exports CSV as cp1252
        return content.decode("cp1252", errors="replace")


def _parse_delimited(content: bytes, delimiter: Optional[str]) -> Tuple[List[str], List[Row]]:
    text = _decode(content)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration:
        return [], []
    columns = _unique_columns(header)
    rows: List[Row] = []
    for line in reader:
        if not any(cell.strip() for cell in line):
            continue
        padded = (line + [""] * len(columns))[: len(columns)]
        # Blank cells are missing values: store NULL (not '') so IS NULL / COUNT / CAST behave
        rows.append({col: (cell if cell.strip() else None) for col, cell in zip(columns, padded)})
    return columns, rows


def _parse_json(content: bytes) -> Tuple[List[str], List[Row]]:
    try:
        data = json.loads(_decode(content))
    except json.JSONDecodeError as e:
        raise UnsupportedFileError(f"Invalid JSON: {e}") from e
    if isinstance(data, dict):
        # {"data": [...]} / {"records": [...]} wrappers, otherwise a single record
        lists = [v for v in data.values() if isinstance(v, list) and v and all(isinstance(i, dict) for i in v)]
        data = lists[0] if len(lists) == 1 else [data]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise UnsupportedFileError("JSON must be an array of objects (or an object wrapping one).")
    return _rows_from_dicts(data)


def _parse_jsonl(content: bytes) -> Tuple[List[str], List[Row]]:
    records: List[Dict[str, Any]] = []
    for lineno, line in enumerate(_decode(content).splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError as e:
            raise UnsupportedFileError(f"Invalid JSON on line {lineno}: {e}") from e
        if not isinstance(rec, dict):
            raise UnsupportedFileError(f"Line {lineno} is not a JSON object.")
        records.append(rec)
    return _rows_from_dicts(records)


def _parse_parquet(content: bytes) -> Tuple[List[str], List[Row]]:
    try:
        import pyarrow.parquet as pq
    except ImportError as e:
        raise UnsupportedFileError("Parquet support requires pyarrow on the server.") from e
    try:
        table = pq.read_table(io.BytesIO(content))
    except Exception as e:  # pyarrow raises several unrelated error types for corrupt files
        raise UnsupportedFileError(f"Could not read Parquet file: {e}") from e
    columns = _unique_columns(table.column_names)
    cols = [table.column(i).to_pylist() for i in range(table.num_columns)]
    rows = [{columns[c]: _stringify(cols[c][r]) for c in range(len(columns))} for r in range(table.num_rows)]
    return columns, rows


def _parse_xlsx(content: bytes) -> Tuple[List[str], List[Row]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise UnsupportedFileError("Excel support requires openpyxl on the server.") from e
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises zipfile/XML errors for corrupt files
        raise UnsupportedFileError(f"Could not read Excel file: {e}") from e
    try:
        sheet = wb.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            return [], []
        columns = _unique_columns(list(header))
        rows: List[Row] = []
        for line in iterator:
            if all(cell is None or str(cell).strip() == "" for cell in line):
                continue
            padded = (list(line) + [None] * len(columns))[: len(columns)]
            rows.append({columns[i]: (_stringify(v) if str(v).strip() else None) if v is not None else None
                         for i, v in enumerate(padded)})
        return columns, rows
    finally:
        wb.close()


def parse_tabular_file(filename: str, content: bytes) -> Tuple[List[str], List[Row]]:
    """Return (columns, rows) for an uploaded file, dispatching on its extension."""
    ext = os.path.splitext(filename or "")[1].lower()
    logger.info("[ingest] stage=parse_file ext=%s bytes=%d", ext, len(content))
    if ext == ".csv":
        return _parse_delimited(content, None)
    if ext == ".tsv":
        return _parse_delimited(content, "\t")
    if ext == ".json":
        return _parse_json(content)
    if ext in (".jsonl", ".ndjson"):
        return _parse_jsonl(content)
    if ext == ".parquet":
        return _parse_parquet(content)
    if ext == ".xlsx":
        return _parse_xlsx(content)
    raise UnsupportedFileError(
        f"Unsupported file type '{ext or filename}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
    )
